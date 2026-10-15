"""生成知识库演示文档《主流大模型（LLM）介绍.docx》。

用法：conda activate dt && python scripts/make_demo_doc.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT_PATH = (
    Path(__file__).resolve().parent.parent
    / "data"
    / "演示文档"
    / "主流大模型（LLM）介绍.docx"
)


def h1(doc, text):
    doc.add_heading(text, level=1)


def h2(doc, text):
    doc.add_heading(text, level=2)


def p(doc, text):
    para = doc.add_paragraph(text)
    para.paragraph_format.space_after = Pt(6)
    return para


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, text in enumerate(header):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    doc.add_paragraph()


def build() -> Document:
    doc = Document()

    title = doc.add_heading("主流大模型（LLM）技术介绍", level=0)

    p(
        doc,
        "本文档是 RAG 系统的本地知识库演示数据，系统性介绍当前主流的大语言模型"
        "（Large Language Model，简称 LLM）：包括发展脉络、代表性模型、关键能力、"
        "主流应用场景、选型建议以及未来趋势，帮助读者快速建立对大模型领域的整体认知。",
    )

    # 一、大模型概述
    h1(doc, "一、大模型概述")
    h2(doc, "1.1 什么是大语言模型")
    p(
        doc,
        "大语言模型（Large Language Model，LLM）是一类基于 Transformer 架构、"
        "在海量文本数据上进行预训练的深度学习模型。它通过对数千亿甚至上万亿 token 的"
        "学习，掌握语言的统计规律与世界知识，从而具备文本生成、理解、推理、翻译、"
        "摘要等多种能力。与传统的规则系统或小模型不同，大模型的参数规模通常在数十亿"
        "到万亿级别，规模越大、数据越多，模型表现出的能力通常越强，这一现象被称为"
        "“规模定律”（Scaling Law）。",
    )
    h2(doc, "1.2 大模型的发展脉络")
    p(
        doc,
        "大模型的发展大致经历了三个阶段。第一阶段是预训练启蒙期（2018—2020 年）："
        "Google 提出 Transformer 架构后，BERT、GPT-1/2 等模型相继问世，验证了"
        "“预训练 + 微调”范式的有效性。第二阶段是规模化突破期（2020—2023 年）："
        "OpenAI 的 GPT-3 以 1750 亿参数展示了上下文学习（In-Context Learning）能力，"
        "随后 RLHF（基于人类反馈的强化学习）技术让模型对齐人类偏好，ChatGPT 的发布"
        "标志着大模型进入大众视野。第三阶段是多模态与推理期（2023 年至今）：模型"
        "开始支持图像、音频、视频等多模态输入，并通过思维链（Chain-of-Thought）与"
        "强化学习显著提升复杂推理能力，同时开源生态快速繁荣。",
    )

    # 二、代表性模型
    h1(doc, "二、代表性模型盘点")
    h2(doc, "2.1 国际主流闭源模型")
    table(
        doc,
        ["厂商", "代表模型", "主要特点"],
        [
            [
                "OpenAI（美国）",
                "GPT 系列（GPT-4o、o 系列推理模型、GPT-5）",
                "大模型浪潮引领者，生态最完善；GPT-4o 支持文本/图像/音频实时多模态交互，"
                "o 系列专注长链推理，数学与代码能力突出。",
            ],
            [
                "Anthropic（美国）",
                "Claude 系列（Claude 3.5/4/4.5 系列）",
                "以安全性和长文本处理见长，支持超长上下文；代码生成与代理（Agent）任务"
                "表现出色，倡导“宪法式 AI”对齐方法。",
            ],
            [
                "Google（美国）",
                "Gemini 系列（Gemini 2.x）",
                "原生多模态设计，文本/图像/音频/视频统一处理；深度整合 Google 搜索、"
                "办公套件等生态，长上下文能力领先。",
            ],
            [
                "xAI（美国）",
                "Grok 系列",
                "主打实时性与幽默风格，深度结合 X（原 Twitter）平台数据。",
            ],
        ],
    )
    h2(doc, "2.2 国内主流模型")
    table(
        doc,
        ["厂商", "代表模型", "主要特点"],
        [
            [
                "DeepSeek（深度求索）",
                "DeepSeek-V3、DeepSeek-R1",
                "以极低训练成本实现顶级性能；R1 专注深度推理，在数学、代码基准上比肩"
                "国际一线模型，并采用完全开源开放权重策略。",
            ],
            [
                "阿里云",
                "通义千问 Qwen 系列（Qwen2.5/Qwen3）",
                "国内最活跃的开源模型家族之一，覆盖 0.5B 到百亿级参数，提供基座、"
                "对话、数学、代码、视觉、嵌入等全系列变体。",
            ],
            [
                "智谱 AI（Zhipu AI）",
                "GLM 系列（GLM-4/4.5/5 系列）",
                "源自清华大学 KEG 实验室技术积累，开源与闭源并行；GLM 系列在通用对话、"
                "代码与智能体任务上表现均衡，并提供免费体验版本。",
            ],
            [
                "月之暗面（Moonshot AI）",
                "Kimi 系列",
                "以超长上下文处理能力著称，早期即支持数十万汉字上下文，擅长长文档"
                "分析、资料整理场景。",
            ],
            [
                "字节跳动",
                "豆包（Doubao）系列",
                "结合字节生态优势，在多模态与大规模 C 端应用中普及度高。",
            ],
            [
                "百度",
                "文心一言（ERNIE 系列）",
                "国内最早发布的大模型产品之一，深耕中文语义理解与产业应用。",
            ],
            [
                "腾讯",
                "混元（Hunyuan）系列",
                "面向腾讯内部业务与腾讯云客户，覆盖文本、图像、视频生成。",
            ],
        ],
    )

    # 三、关键技术
    h1(doc, "三、大模型的关键技术")
    h2(doc, "3.1 预训练与指令微调")
    p(
        doc,
        "大模型的训练通常分为三个阶段：预训练阶段让模型在海量语料上学习语言建模"
        "（预测下一个 token），获得通用知识；监督微调（SFT）阶段使用高质量的"
        "“指令—回答”对数据训练模型遵循人类指令；对齐阶段则通过 RLHF、DPO 等"
        "方法让模型的输出符合人类的价值偏好与安全要求。三阶段训练使模型从“补全"
        "文本的机器”转变为“能够对话协作的助手”。",
    )
    h2(doc, "3.2 注意力机制与上下文窗口")
    p(
        doc,
        "Transformer 的核心是自注意力（Self-Attention）机制，它让模型在处理每个"
        "词时都能关注序列中所有其他词，从而捕捉长距离依赖关系。上下文窗口"
        "（Context Window）指模型一次能处理的最大 token 数量，从早期的 2K—4K "
        "扩展到如今的 128K 乃至百万级别。长上下文使得整本书、整个代码仓库一次性"
        "输入模型成为可能，但也带来注意力稀释与推理成本上升的挑战。",
    )
    h2(doc, "3.3 推理增强与思维链")
    p(
        doc,
        "以 OpenAI o 系列、DeepSeek-R1 为代表的“推理模型”通过思维链"
        "（Chain-of-Thought）让模型在回答前先进行逐步思考，并结合强化学习"
        "（如 GRPO、PPO）训练模型自我检验与纠错，从而在数学证明、竞赛编程、"
        "复杂规划等任务上大幅提升准确率。代价是回答前需要更长的思考时间与"
        "更多的算力消耗。",
    )
    h2(doc, "3.4 多模态与智能体")
    p(
        doc,
        "多模态指模型能够统一理解文本、图像、音频、视频等多种输入形式，"
        "典型应用包括图片理解、视频分析、语音对话等。智能体（Agent）则是"
        "大模型的高阶应用形态：模型通过调用工具（搜索、代码执行、API 调用）、"
        "规划任务步骤、维护长期记忆，自主完成复杂工作流，从“问答机器”进化为"
        "“能够动手执行的数字员工”。",
    )

    # 四、应用场景
    h1(doc, "四、主流应用场景")
    bullets(
        doc,
        [
            "智能问答与客服：基于企业知识库的大模型问答（RAG），提供准确且带引用来源的回答；",
            "内容创作：文案撰写、公文起草、营销创意、剧本与小说辅助创作；",
            "代码助手：代码生成、补全、重构、Bug 修复与代码评审，代表产品如 Copilot、Cursor；",
            "文档处理：合同审查、报告摘要、简历筛选、长文档翻译与信息抽取；",
            "教育培训：个性化辅导、口语陪练、作业批改与学习规划；",
            "科学研究：文献调研、数学证明辅助、蛋白质结构与药物分子设计（AI for Science）；",
            "数据分析：自然语言查询数据库、自动绘制图表与生成分析结论；",
            "智能体应用：自动化办公流程、多系统协同操作、7x24 小时自动运维等。",
        ],
    )

    # 五、选型建议
    h1(doc, "五、模型选型建议")
    p(
        doc,
        "选择大模型时通常需要综合考虑以下维度：任务类型（通用对话、代码、数学推理、"
        "多模态）、性能与准确率、响应速度、调用成本、上下文长度、数据安全要求"
        "（是否支持私有化部署）以及生态与工具链成熟度。一般建议：",
    )
    bullets(
        doc,
        [
            "通用问答与内容生成：可选用 GLM、Qwen、GPT、Claude 等综合能力强的对话模型；",
            "复杂推理与数学代码：优先选择 DeepSeek-R1、o 系列等具备深度推理能力的模型；",
            "本地私有化部署：优先考虑 Qwen、DeepSeek、GLM 等提供开放权重的开源模型；",
            "文本向量化检索（RAG 嵌入）：可选用 Qwen3-Embedding、BGE 等专用嵌入模型；",
            "多模态任务：选择 Gemini、GPT-4o、Qwen-VL 等原生多模态模型。",
        ],
    )

    # 六、发展趋势
    h1(doc, "六、未来发展趋势")
    p(
        doc,
        "第一，推理能力持续深化：模型将在更复杂的数学、科学与工程问题上接近甚至"
        "超过人类专家水平。第二，成本快速下降：训练与推理效率优化（如混合专家模型 "
        "MoE、量化蒸馏）使大模型调用成本以每年数倍的速度降低。第三，智能体规模化"
        "落地：大模型将从“对话”走向“干活”，自主调用工具、跨系统协作成为标配。"
        "第四，多模态走向原生统一：文本、视觉、语音在单一模型中融合处理。第五，"
        "开源与闭源并行发展：开源模型性能不断逼近闭源，为私有化部署与行业定制"
        "提供坚实基础。第六，安全与治理受到重视：各国加快大模型监管立法，"
        "可信 AI 成为产业健康发展的关键。",
    )

    doc.add_paragraph()
    p(doc, "（本文档仅作为 RAG 系统演示数据，内容会随技术发展持续过时，请以各厂商官方发布为准。）")
    return doc


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = build()
    doc.save(OUT_PATH)
    text_len = sum(len(p.text) for p in doc.paragraphs)
    tables = len(doc.tables)
    print(f"已生成：{OUT_PATH}")
    print(f"段落数 {len(doc.paragraphs)}，约 {text_len} 字，表格 {tables} 个")


if __name__ == "__main__":
    main()
